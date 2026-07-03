import io
import re

from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas


def _markdown_lines(markdown_text: str) -> list[tuple[str, str]]:
    """Very small markdown parser: returns (kind, text) where kind is
    heading1 / heading2 / bullet / text / blank.
    """
    lines: list[tuple[str, str]] = []
    for raw_line in (markdown_text or "").splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            lines.append(("blank", ""))
        elif line.startswith("# "):
            lines.append(("heading1", line[2:].strip()))
        elif line.startswith("## "):
            lines.append(("heading2", line[3:].strip()))
        elif line.startswith(("- ", "* ")):
            lines.append(("bullet", line[2:].strip()))
        else:
            lines.append(("text", re.sub(r"[*_`]", "", line).strip()))
    return lines


def markdown_to_pdf_bytes(markdown_text: str) -> bytes:
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=LETTER)
    width, height = LETTER
    x_margin = 0.75 * inch
    y = height - 0.75 * inch
    line_height = 14

    for kind, text in _markdown_lines(markdown_text):
        if y < 0.75 * inch:
            pdf.showPage()
            y = height - 0.75 * inch

        if kind == "blank":
            y -= line_height * 0.6
            continue
        if kind == "heading1":
            pdf.setFont("Helvetica-Bold", 15)
            pdf.drawString(x_margin, y, text)
            y -= line_height * 1.4
            continue
        if kind == "heading2":
            pdf.setFont("Helvetica-Bold", 12)
            pdf.drawString(x_margin, y, text)
            y -= line_height * 1.2
            continue
        if kind == "bullet":
            pdf.setFont("Helvetica", 10.5)
            pdf.drawString(x_margin + 12, y, f"• {text}")
            y -= line_height
            continue

        pdf.setFont("Helvetica", 10.5)
        for wrapped in _wrap(text, 95):
            if y < 0.75 * inch:
                pdf.showPage()
                y = height - 0.75 * inch
            pdf.drawString(x_margin, y, wrapped)
            y -= line_height

    pdf.save()
    return buffer.getvalue()


def markdown_to_docx_bytes(markdown_text: str) -> bytes:
    document = Document()
    for kind, text in _markdown_lines(markdown_text):
        if kind == "blank":
            document.add_paragraph("")
        elif kind == "heading1":
            document.add_heading(text, level=1)
        elif kind == "heading2":
            document.add_heading(text, level=2)
        elif kind == "bullet":
            document.add_paragraph(text, style="List Bullet")
        else:
            document.add_paragraph(text)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _wrap(text: str, width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if len(candidate) > width:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines or [""]
